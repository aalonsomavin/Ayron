import json
import re
from io import BytesIO
from typing import Annotated

from docx import Document
from docx.shared import Pt
from langchain_core.tools import InjectedToolCallId, tool

from apps.agent.context import get_agent_conversation, get_agent_user
from apps.files.services import (
    get_file_for_conversation,
    save_generated_file,
    serialize_file_for_agent,
    serialize_file_for_ui,
    update_generated_file,
)
from apps.files.services import escape_preview_text as esc

_DOCUMENT_DISPLAY_REGISTRY: dict[str, dict] = {}

MAX_SECTIONS = 20
MAX_TABLE_ROWS = 50
MAX_PARAGRAPHS_PER_SECTION = 20
MAX_BULLETS_PER_SECTION = 30

AGENT_INSTRUCTION_AFTER_DOCUMENT = (
    "El documento ya está visible en el chat del usuario. "
    "NO repitas el contenido del informe en tu siguiente mensaje. "
    "Solo añade una frase breve si aporta contexto, o termina sin texto."
)


def pop_document_display(tool_call_id: str | None) -> dict | None:
    if not tool_call_id:
        return None
    return _DOCUMENT_DISPLAY_REGISTRY.pop(tool_call_id, None)


def _sanitize_filename(name: str, fallback: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', "", name).strip()
    if not cleaned:
        cleaned = fallback
    if not cleaned.lower().endswith(".docx"):
        cleaned = f"{cleaned}.docx"
    return cleaned[:200]


def validate_content_json(
    title: str,
    subtitle: str,
    sections: list,
) -> dict:
    if not title or not str(title).strip():
        raise ValueError("title is required")
    if not isinstance(sections, list) or not sections:
        raise ValueError("sections must be a non-empty list")
    if len(sections) > MAX_SECTIONS:
        raise ValueError(f"sections exceeds maximum of {MAX_SECTIONS}")

    normalized_sections = []
    for section in sections:
        if not isinstance(section, dict):
            raise ValueError("each section must be an object")
        heading = str(section.get("heading", "")).strip()
        if not heading:
            raise ValueError("each section requires a heading")
        paragraphs = section.get("paragraphs") or []
        bullets = section.get("bullets") or []
        table = section.get("table")

        if not isinstance(paragraphs, list):
            raise ValueError("paragraphs must be a list")
        if not isinstance(bullets, list):
            raise ValueError("bullets must be a list")
        if len(paragraphs) > MAX_PARAGRAPHS_PER_SECTION:
            raise ValueError(f"paragraphs exceeds maximum of {MAX_PARAGRAPHS_PER_SECTION}")
        if len(bullets) > MAX_BULLETS_PER_SECTION:
            raise ValueError(f"bullets exceeds maximum of {MAX_BULLETS_PER_SECTION}")

        normalized_table = None
        if table is not None:
            if not isinstance(table, dict):
                raise ValueError("table must be an object")
            headers = table.get("headers") or []
            rows = table.get("rows") or []
            if not headers:
                raise ValueError("table requires headers")
            if len(rows) > MAX_TABLE_ROWS:
                raise ValueError(f"table rows exceeds maximum of {MAX_TABLE_ROWS}")
            normalized_table = {
                "headers": [str(h) for h in headers],
                "rows": [[str(cell) for cell in row] for row in rows],
            }

        normalized_sections.append(
            {
                "heading": heading,
                "paragraphs": [str(p) for p in paragraphs],
                "bullets": [str(b) for b in bullets],
                "table": normalized_table,
            }
        )

    return {
        "title": str(title).strip(),
        "subtitle": str(subtitle or "").strip(),
        "sections": normalized_sections,
    }


def build_docx(content_json: dict) -> bytes:
    doc = Document()
    title = content_json.get("title", "")
    subtitle = content_json.get("subtitle", "")

    title_para = doc.add_heading(title, level=0)
    for run in title_para.runs:
        run.font.size = Pt(24)

    if subtitle:
        sub = doc.add_paragraph(subtitle)
        for run in sub.runs:
            run.font.size = Pt(12)

    for section in content_json.get("sections", []):
        doc.add_heading(section.get("heading", ""), level=1)
        for paragraph in section.get("paragraphs", []):
            doc.add_paragraph(paragraph)
        for bullet in section.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")
        table_data = section.get("table")
        if table_data:
            headers = table_data.get("headers", [])
            rows = table_data.get("rows", [])
            if headers:
                table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                table.style = "Table Grid"
                for col_idx, header in enumerate(headers):
                    table.rows[0].cells[col_idx].text = header
                for row_idx, row in enumerate(rows):
                    for col_idx, cell in enumerate(row):
                        if col_idx < len(headers):
                            table.rows[row_idx + 1].cells[col_idx].text = cell

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _render_table_html(table_data: dict) -> str:
    headers = table_data.get("headers", [])
    rows = table_data.get("rows", [])
    if not headers:
        return ""
    head_cells = "".join(f"<th>{esc(h)}</th>" for h in headers)
    body_rows = []
    for row in rows:
        cells = "".join(
            f"<td>{esc(cell)}</td>" for cell in row[: len(headers)]
        )
        while len(row) < len(headers):
            pass
        padded = list(row[: len(headers)])
        while len(padded) < len(headers):
            padded.append("")
        cells = "".join(f"<td>{esc(cell)}</td>" for cell in padded)
        body_rows.append(f"<tr>{cells}</tr>")
    return (
        f'<table class="ay-doc-preview__table"><thead><tr>{head_cells}</tr></thead>'
        f"<tbody>{''.join(body_rows)}</tbody></table>"
    )


def build_preview_html(content_json: dict) -> str:
    title = esc(content_json.get("title", ""))
    subtitle = esc(content_json.get("subtitle", ""))
    parts = [
        '<div class="ay-doc-preview__page">',
        f'<div class="ay-doc-preview__kicker">Documento</div>',
        f'<h1 class="ay-doc-preview__title">{title}</h1>',
    ]
    if subtitle:
        parts.append(f'<div class="ay-doc-preview__subtitle">{subtitle}</div>')
    parts.append('<div class="ay-doc-preview__divider"></div>')

    for section in content_json.get("sections", []):
        heading = esc(section.get("heading", ""))
        parts.append(f'<h2 class="ay-doc-preview__heading">{heading}</h2>')
        for paragraph in section.get("paragraphs", []):
            parts.append(f'<p class="ay-doc-preview__p">{esc(paragraph)}</p>')
        bullets = section.get("bullets", [])
        if bullets:
            items = "".join(f"<li>{esc(b)}</li>" for b in bullets)
            parts.append(f'<ul class="ay-doc-preview__ul">{items}</ul>')
        table_data = section.get("table")
        if table_data:
            parts.append(_render_table_html(table_data))

    parts.append("</div>")
    return "".join(parts)


def _merge_content_json(existing: dict, title, subtitle, sections) -> dict:
    merged = dict(existing)
    if title is not None:
        merged["title"] = str(title).strip()
    if subtitle is not None:
        merged["subtitle"] = str(subtitle).strip()
    if sections is not None:
        merged["sections"] = sections
    return validate_content_json(
        merged.get("title", ""),
        merged.get("subtitle", ""),
        merged.get("sections", []),
    )


def _build_agent_tool_response(file_obj, action: str) -> str:
    return json.dumps(
        {
            "ok": True,
            "action": action,
            "file_id": str(file_obj.id),
            "name": file_obj.original_name,
            "version": file_obj.version,
            "agent_instruction": AGENT_INSTRUCTION_AFTER_DOCUMENT,
        }
    )


def _register_display(tool_call_id: str, file_obj, updated: bool = False) -> None:
    payload = serialize_file_for_ui(file_obj)
    payload["updated"] = updated
    _DOCUMENT_DISPLAY_REGISTRY[tool_call_id] = payload


@tool
def create_document(
    title: str,
    sections: list[dict],
    subtitle: str = "",
    filename: str = "",
    tool_call_id: Annotated[str, InjectedToolCallId] = "",
) -> str:
    """Create a Word document (.docx) for the user in the chat.

    Use when the user asks for a report, memo, or exportable written document.
    Provide structured sections with headings, paragraphs, bullets, and optional tables.
    The document appears in the chat; do not repeat its content in your text response.

    To modify an existing document later, use update_document with the same file_id.
    """
    conversation = get_agent_conversation()
    user = get_agent_user()
    if conversation is None or user is None:
        return json.dumps({"ok": False, "error": "No conversation context"})

    try:
        content_json = validate_content_json(title, subtitle, sections)
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})

    original_name = _sanitize_filename(filename, content_json["title"])
    docx_bytes = build_docx(content_json)
    preview_html = build_preview_html(content_json)

    file_obj = save_generated_file(
        conversation=conversation,
        user=user,
        original_name=original_name,
        content_json=content_json,
        docx_bytes=docx_bytes,
        preview_html=preview_html,
    )
    _register_display(tool_call_id, file_obj, updated=False)
    return _build_agent_tool_response(file_obj, "created")


@tool
def list_conversation_files() -> str:
    """List Word documents generated in the current conversation.

    Returns file_id, name, version, and summary for each document.
    """
    conversation = get_agent_conversation()
    if conversation is None:
        return json.dumps({"ok": False, "error": "No conversation context"})

    from apps.files.services import get_agent_file_index

    files = get_agent_file_index(conversation)
    return json.dumps({"ok": True, "files": files})


@tool
def get_document(file_id: str) -> str:
    """Read the full structured content of a document by file_id.

    Call this before update_document when you need the current sections.
    """
    conversation = get_agent_conversation()
    if conversation is None:
        return json.dumps({"ok": False, "error": "No conversation context"})

    file_obj = get_file_for_conversation(file_id, conversation)
    if file_obj is None:
        return json.dumps({"ok": False, "error": "Document not found in this conversation"})

    return json.dumps(
        {
            "ok": True,
            **serialize_file_for_agent(file_obj),
            "content_json": file_obj.content_json,
        }
    )


@tool
def update_document(
    file_id: str,
    title: str | None = None,
    subtitle: str | None = None,
    sections: list[dict] | None = None,
    tool_call_id: Annotated[str, InjectedToolCallId] = "",
) -> str:
    """Update an existing Word document by file_id.

    Provide only the fields you want to change. sections replaces all sections when provided.
    """
    conversation = get_agent_conversation()
    if conversation is None:
        return json.dumps({"ok": False, "error": "No conversation context"})

    file_obj = get_file_for_conversation(file_id, conversation)
    if file_obj is None:
        return json.dumps({"ok": False, "error": "Document not found in this conversation"})

    try:
        content_json = _merge_content_json(file_obj.content_json, title, subtitle, sections)
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})

    docx_bytes = build_docx(content_json)
    preview_html = build_preview_html(content_json)
    file_obj = update_generated_file(
        file_obj=file_obj,
        content_json=content_json,
        docx_bytes=docx_bytes,
        preview_html=preview_html,
    )
    _register_display(tool_call_id, file_obj, updated=True)
    return _build_agent_tool_response(file_obj, "updated")
